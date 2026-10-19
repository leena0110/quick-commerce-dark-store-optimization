import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Helper function to add title and content slide
def add_slide(title, content_list):
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    for i, point in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
            p.text = point
        else:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
    return slide

# 1. Title Slide
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Quick-Commerce Dark Store Inventory Optimization"
subtitle.text = "Review 1 - Unit 1 & 2\n23CSE452 — Business Analytics Capstone Project\n\nTeam Members:\nLEENA SRI K (CB.SC.U4CSE23526)\nA.C SANHITHA REDDY (CB.SC.U4CSE23206)\nKANISH KUMARAN M (CB.SC.U4CSE23424)\nANDE TARAKA SRINIVAS (CB.SC.U4CSE23212)\nAJAY N MUKUNTH (CB.SC.U4CSE23006)"

# 2. Problem Statement & Business Objective
add_slide("Problem Statement & Business Objective", [
    "Context: Quick-commerce platforms (Blinkit, Zepto, Swiggy) rely on dark stores.",
    "Problem: High inventory holding costs vs stockouts. Need demand-informed inventory planning.",
    "Business Objective (Review 1): Predict daily product demand at the dark-store level to support inventory planning.",
    "Future Scope (Review 2): Identify frequently co-purchased products for product bundling."
])

# 3. Dataset Description
add_slide("Dataset Description & Generation", [
    "Method: 100% Synthetic data generated using Python (NumPy/Pandas) for 8 stores.",
    "Size: ~106,000 raw transaction records (April - Sept 2026).",
    "Scope: 58 products across 9 categories.",
    "Assumptions: Includes realistic hourly, daily, and seasonal demand variations.",
    "Features: Order_ID, Timestamp, Customer_ID, Quantity, Unit_Price, Discount, etc."
])

# 4. Data Cleaning & Preprocessing
add_slide("Data Cleaning & Preprocessing", [
    "Total Original Records: 106,398",
    "Cleaned Records: 105,872 (526 duplicates removed)",
    "Missing Values: Delivery Distance imputed with store-level medians.",
    "Anomalies: Corrected negative prices (absolute value) and capped extremely high quantities.",
    "Final Dataset: Reliable, quality-assured dataset ready for modeling."
])

# 5. Exploratory Data Analysis (EDA)
add_slide("Exploratory Data Analysis", [
    "Peak-Hour Demand: 52.2% of sales occur in morning (7-9 AM) and evening (6-9 PM) peaks.",
    "Top Categories: Fruits & Vegetables and Dairy represent the highest volume.",
    "Weekend vs Weekday: Average demand is 23.1% higher on weekends.",
    "These insights drive feature selection and baseline inventory assumptions."
])

# 6. Feature Engineering
add_slide("Feature Engineering", [
    "Transformed transaction data into daily SKU-store demand panel.",
    "Target Variable: Daily_Quantity (Regression).",
    "Calendar Features: Day of Week, Month, Is_Weekend.",
    "Lag & Rolling Features: Demand_Lag_1d, Demand_Roll_Mean_7d.",
    "Historical Averages: Product and Category historical average demand.",
    "Total of 18 leakage-safe features developed."
])

# 7. Predictive Modelling
add_slide("Predictive Modelling", [
    "Approach: Regression task using chronological 80/20 train-test split.",
    "Models Developed: Linear Regression, Random Forest Regressor, XGBoost Regressor.",
    "Evaluation Metrics: R², MAE, MSE, RMSE.",
    "Best Model: Random Forest Regressor (R² = 0.3710, MAE = 1.3322, RMSE = 1.8332).",
    "Key Insight: Historical product demand is the most important feature (65% importance)."
])

# 8. Individual Contribution & Tracking
add_slide("Individual Contribution Summary", [
    "LEENA SRI K: Project Management (GitHub Projects setup), Dataset generation & documentation.",
    "A.C SANHITHA REDDY: Data cleaning, missing value imputation, and quality validation.",
    "KANISH KUMARAN M: Exploratory Data Analysis (EDA) and visualizations.",
    "ANDE TARAKA SRINIVAS: Feature Engineering (Lag, Rolling, and Calendar variables).",
    "AJAY N MUKUNTH: Predictive Modeling, hyperparameter tuning, and model evaluation."
])

# 9. Project Tracking
add_slide("Project Tracking", [
    "Project Management Tool: GitHub Projects",
    "Board Columns: Backlog, To Do, In Progress, Review/Testing, Completed.",
    "All tasks are actively tracked and updated by assigned team members.",
    "Link to Board: https://github.com/users/leena0110/projects/1 (Placeholder link)"
])

# Save presentation
prs.save('Review_1_Presentation_Format.pptx')
print("Presentation generated successfully at Review_1_Presentation_Format.pptx")
