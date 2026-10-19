from pptx import Presentation

# Use the original template (not our already-modified output)
prs = Presentation(r'Final_Project_Presentation_Template.pptx')
for i, slide in enumerate(prs.slides):
    print(f'\n=== SLIDE {i+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if '[' in txt or (len(txt) > 5 and txt not in ['', '\x0b']):
                print(repr(txt[:400]))
